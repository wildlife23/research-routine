import sys
import json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

def create_weekly_review(content_json):
    data = json.loads(content_json)
    doc = Document()

    # 기본 스타일 설정
    style = doc.styles['Normal']
    style.font.name = '맑은 고딕'
    style.font.size = Pt(10)

    # 표지
    title = doc.add_heading('', 0)
    title_run = title.add_run('주간 논문 리뷰 보고서')
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(0x1a, 0x56, 0x76)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_para.add_run(data.get('date', ''))
    date_run.font.size = Pt(12)
    date_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph()

    # 요약 통계 박스
    doc.add_heading('이번 주 요약', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    cells = table.rows[0].cells
    cells[0].text = f"검색 논문 수\n{data.get('total_papers', 0)}편"
    cells[1].text = f"커버 분야\n{data.get('total_fields', 0)}개"
    cells[2].text = f"주목 논문\n{len(data.get('top3', []))}편"
    for cell in cells:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # TOP 3
    doc.add_heading('이번 주 주목 논문 TOP 3', level=1)
    for i, paper in enumerate(data.get('top3', []), 1):
        p = doc.add_paragraph()
        run = p.add_run(f"{i}. {paper['title']}")
        run.bold = True
        run.font.color.rgb = RGBColor(0x1a, 0x56, 0x76)
        doc.add_paragraph(paper.get('reason', ''))
        doc.add_paragraph()

    doc.add_page_break()

    # 분야별 논문 요약
    doc.add_heading('분야별 논문 요약', level=1)
    for field in data.get('fields', []):
        doc.add_heading(field['name'], level=2)
        for paper in field.get('papers', []):
            doc.add_heading(paper['title'], level=3)
            table = doc.add_table(rows=6, cols=2)
            table.style = 'Table Grid'
            rows_data = [
                ('저자 / 저널 / 연도', paper.get('meta', '')),
                ('연구 목적', paper.get('objective', '')),
                ('핵심 방법론', paper.get('method', '')),
                ('주요 결과', paper.get('results', '')),
                ('보전·복원 시사점', paper.get('implications', '')),
                ('내 연구와의 연관성', paper.get('relevance', ''))
            ]
            for i, (label, value) in enumerate(rows_data):
                row = table.rows[i]
                row.cells[0].text = label
                row.cells[1].text = value
                row.cells[0].paragraphs[0].runs[0].bold = True
            doc.add_paragraph()

    doc.add_page_break()

    # Zotero 등록용 표
    doc.add_heading('Zotero 등록용 정보', level=1)
    zotero_table = doc.add_table(rows=1, cols=5)
    zotero_table.style = 'Table Grid'
    headers = ['제목', '저자', '저널', '연도', 'DOI']
    for i, h in enumerate(headers):
        cell = zotero_table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for entry in data.get('zotero', []):
        row = zotero_table.add_row()
        row.cells[0].text = entry.get('title', '')
        row.cells[1].text = entry.get('author', '')
        row.cells[2].text = entry.get('journal', '')
        row.cells[3].text = entry.get('year', '')
        row.cells[4].text = entry.get('doi', '')

    filename = f"{data.get('date', 'weekly-review')}-weekly-review.docx"
    doc.save(filename)
    print(f"saved:{filename}")

if __name__ == '__main__':
    create_weekly_review(sys.argv[1])
